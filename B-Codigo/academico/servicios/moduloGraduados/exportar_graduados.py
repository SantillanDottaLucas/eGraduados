
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def generar_excel_graduados(graduados):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Graduados"

    encabezados = [
        "Legajo",
        "DNI",
        "Nombre",
        "Apellido",
        "Fecha de ingreso a la universidad",
        "Carrera",
        "Fecha de ingreso a la carrera",
        "Fecha de egreso de la carrera",
        "Promedio sin aplazo",
        "Promedio con aplazo",
    ]

    hoja.append(encabezados)

    for graduado in graduados.iterator(chunk_size=500):
        carreras = graduado.carreras_cursadas.select_related(
            "codigo_carrera"
        ).all()

        for graduado_carrera in carreras:
            hoja.append([
                graduado.legajo_graduado,
                graduado.dni,
                graduado.nombre,
                graduado.apellido,
                graduado.fecha_ingreso_universidad,
                graduado_carrera.codigo_carrera.nombre_carrera,
                graduado_carrera.fecha_ingreso_carrera,
                graduado_carrera.fecha_egreso_carrera,
                float(graduado_carrera.promedio_sin_aplazo)
                if graduado_carrera.promedio_sin_aplazo is not None
                else None,
                float(graduado_carrera.promedio_con_aplazo)
                if graduado_carrera.promedio_con_aplazo is not None
                else None,
            ])

    # Formato de los encabezados
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill(
            fill_type="solid",
            fgColor="244062"
        )
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    # Formato de fechas y promedios
    for fila in hoja.iter_rows(min_row=2):
        for indice in (5, 7, 8):
            fila[indice - 1].number_format = "dd/mm/yyyy"

        for indice in (9, 10):
            fila[indice - 1].number_format = "0.00"

    # Ancho de columnas
    anchos = [22, 14, 22, 22, 25, 42, 24, 24, 20, 20]

    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho

    hoja.row_dimensions[1].height = 35
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    # Convertir los datos en una tabla de Excel
    if hoja.max_row >= 2:
        tabla = Table(
            displayName="TablaGraduados",
            ref=hoja.dimensions
        )

        estilo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )

        tabla.tableStyleInfo = estilo
        hoja.add_table(tabla)

    archivo = BytesIO()
    libro.save(archivo)

    return archivo.getvalue()
